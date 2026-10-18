"""Generate Attendance-base.xlsx: 15 fictional employees x 30 days (codes P/A/L/H).

Deterministic (seed 42). Writes the sample data AND a stats JSON with the
independently computed expected values used by the proof in prove.py.
100% local. by Muhammad Bilal.
"""
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(HERE, "Attendance-base.xlsx")
OUT_JSON = os.path.join(HERE, "expected-stats.json")

EMPLOYEES = [
    "Ayesha Khan", "Ahmed Raza", "Fatima Noor", "Usman Tariq", "Sana Malik",
    "Imran Sheikh", "Hina Aslam", "Kamran Iqbal", "Nadia Hussain",
    "Faisal Mahmood", "Rabia Anwar", "Shahid Latif", "Maria Siddiqui",
    "Tariq Aziz", "Saima Butt",
]
DAYS = 30
# code distribution: P 88%, A 5%, L 4%, H 3%
CODES = ["P"] * 88 + ["A"] * 5 + ["L"] * 4 + ["H"] * 3

HDR_FILL = PatternFill("solid", fgColor="0E4C81")
HDR_FONT = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center")


def main():
    rng = random.Random(42)
    grid = []
    for _ in EMPLOYEES:
        grid.append([rng.choice(CODES) for _ in range(DAYS)])
    # make the demo illustrative: one chronic absentee (>=5 -> red flag),
    # one frequent latecomer (>=3 -> amber flag)
    grid[7][0:6] = ["A"] * 6      # Kamran Iqbal: 6 absents
    grid[3][10:14] = ["L"] * 4    # Usman Tariq: 4 lates

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.cell(row=1, column=1, value="Employee").font = HDR_FONT
    ws.cell(row=1, column=1).fill = HDR_FILL
    ws.column_dimensions["A"].width = 20
    for d in range(1, DAYS + 1):
        c = ws.cell(row=1, column=1 + d, value=f"Day {d}")
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = CENTER
        ws.column_dimensions[chr(64 + 1 + d) if 1 + d <= 26 else f"A{chr(64 + 1 + d - 26)}"].width = 7
    # fix column letters beyond Z properly
    from openpyxl.utils import get_column_letter
    for d in range(1, DAYS + 1):
        ws.column_dimensions[get_column_letter(1 + d)].width = 7
    for r, emp in enumerate(EMPLOYEES, start=2):
        ws.cell(row=r, column=1, value=emp)
        for d, code in enumerate(grid[r - 2], start=1):
            c = ws.cell(row=r, column=1 + d, value=code)
            c.alignment = CENTER

    # independent expected stats (mirror of the macro's counting rules)
    per_emp = []
    tot_absent = 0
    tot_late = 0
    max_absent = -1
    max_absent_emp = ""
    for emp, row in zip(EMPLOYEES, grid):
        p = row.count("P")
        a = row.count("A")
        l = row.count("L")
        h = row.count("H")
        tracked = len(row)
        per_emp.append({
            "employee": emp, "present": p, "absent": a, "late": l,
            "half": h, "tracked": tracked,
            "pct": round(p / tracked * 100, 1),
        })
        tot_absent += a
        tot_late += l
        if a > max_absent:
            max_absent = a
            max_absent_emp = emp
    expected = {
        "employees": len(EMPLOYEES),
        "total_absent": tot_absent,
        "total_late": tot_late,
        "max_absent": max_absent,
        "max_absent_emp": max_absent_emp,
        "per_employee": per_emp,
    }
    wb.save(OUT_XLSX)
    with open(OUT_JSON, "w") as f:
        json.dump(expected, f, indent=2)
    print(f"Wrote {OUT_XLSX} ({len(EMPLOYEES)} employees x {DAYS} days)")
    print(f"Expected: total_absent={tot_absent} total_late={tot_late} "
          f"max_absent={max_absent} ({max_absent_emp})")


if __name__ == "__main__":
    main()
